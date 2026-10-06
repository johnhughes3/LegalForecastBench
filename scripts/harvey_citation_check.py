"""Reproduce the inputs to, and publish the results of, the LAB citation check.

The check read every supplied document and the ``task.json`` of all 52 Harvey LAB
litigation-dispute-resolution tasks, listed every citation to legal authority, checked
each one against CourtListener or another reputable source, and gave every problem
finding an independent second check. The method and results are described in
``docs/harvey-lab-audit/litigation-dispute-resolution/citation-check/README.md``.

Subcommands:

``extract``
    Convert a checkout of ``harveyai/harvey-labs`` at the pinned commit into the plain
    text the checking agents read: DOCX through pandoc (tracked changes shown), XLSX
    cell by cell with comments, EML through Python's ``email`` module, PPTX slide text
    (needs ``--with python-pptx``),
    and ``task.json`` pretty-printed. Needs ``pandoc``.

``publish``
    Write the per-task results, ``problems.md``, and ``summary.json`` from the
    combined result file the checking runs produced.

Examples::

    uv run --with python-pptx python scripts/harvey_citation_check.py extract \\
        --harvey-labs /path/to/harvey-labs --out /path/to/text
    uv run python scripts/harvey_citation_check.py publish --state /path/to/state.json
"""

from __future__ import annotations

import argparse
import email
import importlib
import json
import re
import subprocess
from collections import Counter
from email import policy
from pathlib import Path
from typing import Any

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
CHECK_DIR = ROOT / "docs/harvey-lab-audit/litigation-dispute-resolution/citation-check"
HARVEY_COMMIT = "1dd81403b2fbb60596f7aea3fcecafad7bf73143"
TASKS = "tasks/litigation-dispute-resolution"
PROBLEMS = (
    "NOT_FOUND_LIKELY_FABRICATED",
    "CITATION_POINTS_TO_DIFFERENT_CASE",
    "MISCHARACTERIZED",
    "MISQUOTED",
    "WRONG_CITATION_REAL_CASE",
)
ORDER = ("VERIFIED", "VERIFIED_MINOR_DEFECT", *PROBLEMS, "UNABLE_TO_VERIFY")
TYPES = (
    "case",
    "statute",
    "regulation",
    "court_rule",
    "restatement_or_treatise",
    "other",
)
# Working paths recorded by the checking agents; only the file name is published.
LOCAL_PATH = re.compile(r"(?:/[\w.\-]+)+/text/[\w.\-]+/")


def _docx(path: Path) -> str:
    args = ["pandoc", "-t", "gfm", "--wrap=none", str(path)]
    if path.suffix == ".docx":
        args.insert(1, "--track-changes=all")
    return subprocess.run(args, capture_output=True, text=True, check=True).stdout


def _pptx(path: Path) -> str:
    try:
        pptx: Any = importlib.import_module("pptx")
    except ImportError as exc:
        raise SystemExit(
            "PPTX files need python-pptx: uv run --with python-pptx ..."
        ) from exc
    out: list[str] = []
    for number, slide in enumerate(pptx.Presentation(path).slides, 1):
        out.append(f"## Slide {number}")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                out += [
                    " | ".join(c.text for c in row.cells) for row in shape.table.rows
                ]
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        if notes.strip():
            out.append("[notes] " + notes)
    return "\n".join(out)


def _xlsx(path: Path) -> str:
    book = openpyxl.load_workbook(path, data_only=False)
    out: list[str] = []
    for sheet in book.worksheets:
        out.append(f"## Sheet: {sheet.title}")
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    out.append(f"{cell.coordinate}: {cell.value}")
                if cell.comment:
                    out.append(f"{cell.coordinate} [comment]: {cell.comment.text}")
    return "\n".join(out)


def _eml(path: Path) -> str:
    message: Any = email.message_from_bytes(path.read_bytes(), policy=policy.default)
    out = [
        f"{h}: {message[h]}"
        for h in ("From", "To", "Cc", "Date", "Subject")
        if message[h]
    ]
    for part in message.walk():
        if part.get_content_maintype() == "text":
            out.append(
                f"\n--- part {part.get_content_type()} ---\n{part.get_content()}"
            )
        elif part.get_filename():
            out.append(f"\n--- attachment {part.get_filename()} (not extracted) ---")
    return "\n".join(out)


CONVERTERS = {".docx": _docx, ".xlsx": _xlsx, ".eml": _eml, ".pptx": _pptx}


def extract(harvey_labs: Path, out: Path) -> None:
    """Write one text file per supplied document and per ``task.json``."""
    tasks = sorted(p for p in (harvey_labs / TASKS).iterdir() if p.is_dir())
    for task in tasks:
        dest = out / task.name
        dest.mkdir(parents=True, exist_ok=True)
        spec = json.loads((task / "task.json").read_text(encoding="utf-8"))
        (dest / "task.json.md").write_text(
            json.dumps(spec, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        documents = task / "documents"
        for path in sorted(documents.rglob("*")) if documents.exists() else []:
            if not path.is_file():
                continue
            convert = CONVERTERS.get(path.suffix.lower())
            text = (
                convert(path)
                if convert
                else path.read_text(encoding="utf-8", errors="replace")
            )
            name = path.relative_to(documents).as_posix().replace("/", "__") + ".md"
            (dest / name).write_text(text, encoding="utf-8")
    print(f"wrote text for {len(tasks)} tasks to {out}")


def final_status(citation: dict[str, Any]) -> str:
    """The status after the second check, or the first check's if none was needed."""
    second = citation.get("second_check")
    first: str = citation["first_check"]["status"]
    return str(second["revised_status"]) if second else first


def _scrub(value: Any) -> Any:
    if isinstance(value, str):
        return LOCAL_PATH.sub("", value)
    if isinstance(value, list):
        return [_scrub(v) for v in value]  # pyright: ignore[reportUnknownVariableType]
    if isinstance(value, dict):
        return {k: _scrub(v) for k, v in value.items()}  # pyright: ignore[reportUnknownVariableType]
    return value


def _citation(raw: dict[str, Any]) -> dict[str, Any]:
    out = {
        "ref": raw["ref"],
        "document": raw["document"],
        "location": raw["location"],
        "authority_type": raw["authority_type"],
        "authority": raw["authority_name"],
        "citation_as_written": raw["citation_as_written"],
        "pincite": raw["pincite"],
        "cited_for": raw["proposition"],
        "quotes": raw["quotes"],
        "context": raw["context_excerpt"],
        "found_by": raw["found_by"],
        "fictional_by_design": raw["fictional_by_design"],
        "fictional_reason": raw["fictional_reason"],
    }
    if raw.get("verification"):
        first = dict(raw["verification"])
        first.pop("ref", None)
        out["first_check"] = first
    if raw.get("refutation"):
        second = dict(raw["refutation"])
        second.pop("ref", None)
        out["second_check"] = second
    return _scrub(out)


def _cell(text: str, limit: int = 0) -> str:
    text = " ".join(str(text).split()).replace("|", "\\|")
    return text if not limit or len(text) <= limit else text[: limit - 1] + "…"


def publish(state_path: Path) -> None:
    """Write results/, problems.md, and summary.json from the combined state file."""
    state: dict[str, dict[str, Any]] = json.loads(
        state_path.read_text(encoding="utf-8")
    )
    if len(state) != 52:
        raise SystemExit(f"expected 52 tasks, found {len(state)}")
    results = CHECK_DIR / "results"
    results.mkdir(parents=True, exist_ok=True)
    by_status: Counter[str] = Counter()
    by_type: Counter[tuple[str, str]] = Counter()
    second: Counter[tuple[str, str]] = Counter()
    where: Counter[tuple[str, str]] = Counter()
    rows: list[str] = []
    per_task: list[dict[str, Any]] = []
    instances = fictional = 0
    for task in sorted(state):
        citations = [_citation(c) for c in state[task]["citations"]]
        real = [c for c in citations if not c["fictional_by_design"]]
        if any("first_check" not in c for c in real):
            raise SystemExit(f"{task}: a citation has no first check")
        for citation in real:
            status = final_status(citation)
            by_status[status] += 1
            by_type[(citation["authority_type"], status)] += 1
            source = (
                "task.json"
                if citation["document"].startswith("task.json")
                else "documents"
            )
            where[(source, "problem" if status in PROBLEMS else "other")] += 1
            if citation["first_check"]["status"] in PROBLEMS:
                if "second_check" not in citation:
                    raise SystemExit(f"{task}: a problem finding has no second check")
                held = "overturned" if citation["second_check"]["refuted"] else "held"
                second[(citation["first_check"]["status"], held)] += 1
        instances += len(citations)
        fictional += len(citations) - len(real)
        (results / f"{task}.json").write_text(
            json.dumps(
                {
                    "task": task,
                    "harvey_commit": HARVEY_COMMIT,
                    "documents_read": sorted(
                        {
                            Path(f).name
                            for b in state[task]["batches"]
                            for f in b["files_read"]
                        }
                    ),
                    "citations": citations,
                },
                indent=2,
                ensure_ascii=False,
            )
            + "\n",
            encoding="utf-8",
        )
        problems = [c for c in real if final_status(c) in PROBLEMS]
        per_task.append({"task": task, "checked": len(real), "problems": len(problems)})
        for c in sorted(
            problems, key=lambda c: (PROBLEMS.index(final_status(c)), c["ref"])
        ):
            explanation = (c.get("second_check") or c["first_check"])["explanation"]
            rows.append(
                f"| {task} | {_cell(c['document'])} | {_cell(c['location'], 60)} | "
                f"{_cell(c['authority'], 70)} {_cell(c['citation_as_written'], 60)} | "
                f"{final_status(c)} | {_cell(explanation, 280)} |"
            )
    summary = {
        "harvey_commit": HARVEY_COMMIT,
        "tasks": len(state),
        "citation_instances": instances,
        "fictional_by_design": fictional,
        "checked": sum(by_status.values()),
        "by_final_status": {s: by_status[s] for s in ORDER},
        "problems": sum(by_status[s] for s in PROBLEMS),
        "tasks_with_problems": sum(1 for t in per_task if t["problems"]),
        "by_authority_type": {
            t: {s: by_type[(t, s)] for s in ORDER if by_type[(t, s)]} for t in TYPES
        },
        "second_check": {
            s: {"held": second[(s, "held")], "overturned": second[(s, "overturned")]}
            for s in PROBLEMS
        },
        "by_location": {
            "task_json": {k: where[("task.json", k)] for k in ("problem", "other")},
            "supplied_documents": {
                k: where[("documents", k)] for k in ("problem", "other")
            },
        },
        "per_task": per_task,
    }
    (CHECK_DIR / "summary.json").write_text(
        json.dumps(summary, indent=2) + "\n", encoding="utf-8"
    )
    (CHECK_DIR / "problems.md").write_text(
        "# Problem citations in the LAB litigation task environments\n\n"
        "Every citation whose final status, after the independent second check, is a "
        "problem. AI-generated findings; see the [method and caveats](README.md). Full "
        "evidence for each row, including the source URLs and excerpts the checking "
        "agents read, is in [`results/`](results/).\n\n"
        "| Task | Document | Location | Authority as written | Final status | "
        "Finding |\n|---|---|---|---|---|---|\n" + "\n".join(rows) + "\n",
        encoding="utf-8",
    )
    print(
        json.dumps(
            {k: summary[k] for k in ("checked", "problems", "tasks_with_problems")}
        )
    )


def main() -> None:
    parser = argparse.ArgumentParser(description=(__doc__ or "").split("\n\n")[0])
    sub = parser.add_subparsers(dest="command", required=True)
    one = sub.add_parser(
        "extract", help="convert task files to the text the agents read"
    )
    one.add_argument("--harvey-labs", type=Path, required=True)
    one.add_argument("--out", type=Path, required=True)
    two = sub.add_parser(
        "publish", help="write the published results from a state file"
    )
    two.add_argument("--state", type=Path, required=True)
    args = parser.parse_args()
    if args.command == "extract":
        extract(args.harvey_labs, args.out)
    else:
        publish(args.state)


if __name__ == "__main__":
    main()
